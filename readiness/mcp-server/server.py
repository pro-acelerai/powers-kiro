"""
MCP Server - Readiness
Fornece tools para leitura de historias em Markdown, HTML e PDF, calculo
deterministico dos indices de Definition of Ready / AI-Ready e exportacao
do relatorio de analise.
Protocolo: stdio (JSON-RPC via stdin/stdout)
"""

import json
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import openpyxl
except ImportError:
    openpyxl = None


# === Extensoes suportadas ===

EXT_TEXTO = {".md", ".markdown", ".txt", ".text"}
EXT_HTML = {".html", ".htm", ".xhtml"}
EXT_PDF = {".pdf"}
EXT_SUPORTADAS = EXT_TEXTO | EXT_HTML | EXT_PDF

LIMITE_CARACTERES = 400_000


# === Leitura: Markdown / texto ===

def ler_texto_simples(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


# === Leitura: HTML ===

class ExtratorHTML(HTMLParser):
    """Extrai texto legivel de HTML preservando estrutura de blocos e tabelas."""

    IGNORAR = {"script", "style", "noscript", "head", "meta", "link", "svg"}
    BLOCOS = {
        "p", "div", "section", "article", "header", "footer", "br", "hr",
        "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "table", "thead",
        "tbody", "blockquote", "pre", "ul", "ol", "dl", "dt", "dd",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.partes = []
        self._pilha_ignorada = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.IGNORAR:
            self._pilha_ignorada += 1
            return
        if tag in ("td", "th"):
            self.partes.append(" | ")
        elif tag in self.BLOCOS:
            self.partes.append("\n")

    def handle_endtag(self, tag):
        if tag in self.IGNORAR:
            self._pilha_ignorada = max(0, self._pilha_ignorada - 1)
            return
        if tag in self.BLOCOS:
            self.partes.append("\n")

    def handle_data(self, data):
        if self._pilha_ignorada:
            return
        texto = data.strip()
        if texto:
            self.partes.append(re.sub(r"\s+", " ", texto))

    def resultado(self) -> str:
        bruto = "".join(self.partes)
        bruto = re.sub(r"[ \t]*\n[ \t]*", "\n", bruto)
        bruto = re.sub(r"\n{3,}", "\n\n", bruto)
        linhas = [linha.strip(" |").strip() if linha.strip() == "|" else linha.strip()
                  for linha in bruto.split("\n")]
        return "\n".join(linha for linha in linhas if linha).strip()


def ler_html(path: Path) -> str:
    parser = ExtratorHTML()
    parser.feed(ler_texto_simples(path))
    parser.close()
    return parser.resultado()


# === Leitura: PDF ===

def ler_pdf(path: Path) -> str:
    if pypdf is None:
        return ("ERRO: dependencia 'pypdf' nao instalada. "
                "Execute: pip install -r mcp-server/requirements.txt")

    reader = pypdf.PdfReader(str(path))
    if getattr(reader, "is_encrypted", False):
        try:
            if not reader.decrypt(""):
                raise ValueError("senha necessaria")
        except Exception:
            return "ERRO: PDF protegido por senha. Forneca uma versao sem protecao."

    paginas = []
    paginas_com_texto = 0
    for indice, pagina in enumerate(reader.pages, 1):
        try:
            texto = pagina.extract_text() or ""
        except Exception as exc:
            texto = f"[falha ao extrair texto: {exc}]"
        texto = re.sub(r"\n{3,}", "\n\n", texto.strip())
        if len(re.sub(r"\s", "", texto)) >= 3:
            paginas_com_texto += 1
            paginas.append(f"--- PAGINA {indice} ---\n{texto}")
        else:
            paginas.append(f"--- PAGINA {indice} --- (sem texto extraivel)")

    if paginas_com_texto == 0:
        return ("AVISO: nenhum texto extraivel encontrado no PDF "
                f"({len(paginas)} pagina(s) verificada(s)). "
                "Provavelmente e um PDF digitalizado (imagem). "
                "Solicite ao usuario a versao em texto (.md/.html) ou um PDF pesquisavel.")

    return "\n\n".join(paginas).strip()


# === Dispatcher de leitura ===

def ler_documento(caminho: str) -> str:
    path = Path(caminho).expanduser()
    if not path.exists():
        raise FileNotFoundError(f"Arquivo nao encontrado: {caminho}")
    if path.is_dir():
        raise IsADirectoryError(
            f"'{caminho}' e um diretorio. Use a tool 'listar_documentos' para inspecionar a pasta.")

    ext = path.suffix.lower()
    if ext in EXT_TEXTO:
        tipo, conteudo = "MARKDOWN/TEXTO", ler_texto_simples(path)
    elif ext in EXT_HTML:
        tipo, conteudo = "HTML", ler_html(path)
    elif ext in EXT_PDF:
        tipo, conteudo = "PDF", ler_pdf(path)
    else:
        raise ValueError(
            f"Extensao '{ext}' nao suportada. Suportadas: "
            f"{', '.join(sorted(EXT_SUPORTADAS))}")

    truncado = ""
    if len(conteudo) > LIMITE_CARACTERES:
        conteudo = conteudo[:LIMITE_CARACTERES]
        truncado = (f"\n\n[AVISO: conteudo truncado em {LIMITE_CARACTERES} caracteres. "
                    "Documento muito grande - valide as historias em lotes.]")

    cabecalho = (f"=== DOCUMENTO: {path.name} ===\n"
                 f"Tipo: {tipo}\n"
                 f"Caminho: {path.absolute()}\n"
                 f"Tamanho extraido: {len(conteudo)} caracteres\n"
                 f"{'=' * 60}\n")
    return cabecalho + conteudo + truncado


# === Descoberta de documentos no projeto ===

# Padroes flexiveis que reconhecem diferentes convencoes de ID de historia:
# HU/HT (convencao PT-BR), US (User Story), STORY, F (Feature), etc.
PISTAS_HISTORIA = [
    (re.compile(r"\b(HU|HT|US|STORY|FEAT|F)\s*-?\s*\d{2,}", re.IGNORECASE), "IDs de historia"),
    (re.compile(r"(eu[,\s]+como|as\s+a|como\s+um)\b", re.IGNORECASE), "narrativa de historia"),
    (re.compile(r"crit[eé]rios?\s+de\s+aceite|acceptance\s+criteria", re.IGNORECASE), "criterios de aceite"),
    (re.compile(r"regras?\s+de\s+neg[oó]cio|business\s+rules?", re.IGNORECASE), "regras de negocio"),
    (re.compile(r"\bRN-?\d{2,}", re.IGNORECASE), "IDs de regra (RN)"),
    (re.compile(r"hist[oó]rias?\s+de\s+usu[aá]rio|user\s+stor(y|ies)", re.IGNORECASE), "titulo de historias"),
    (re.compile(r"\b(dado que|when|given)\b", re.IGNORECASE), "cenarios BDD"),
]

IGNORAR_DIRS = {".git", "node_modules", ".venv", "venv", "__pycache__", "dist",
                "build", ".idea", ".vscode", "target", ".next", "coverage"}


def listar_documentos(diretorio: str, profundidade: int = 4) -> str:
    base = Path(diretorio).expanduser()
    if not base.exists():
        raise FileNotFoundError(f"Diretorio nao encontrado: {diretorio}")
    if not base.is_dir():
        raise NotADirectoryError(f"'{diretorio}' nao e um diretorio.")

    encontrados = []
    for path in base.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in EXT_SUPORTADAS:
            continue
        rel = path.relative_to(base)
        if any(parte in IGNORAR_DIRS for parte in rel.parts):
            continue
        if len(rel.parts) - 1 > profundidade:
            continue
        encontrados.append((rel, path))

    if not encontrados:
        return (f"Nenhum documento .md/.html/.pdf encontrado em: {base.absolute()}\n"
                "Peca ao usuario para colar as historias no chat ou indicar o caminho do arquivo.")

    linhas = [f"=== DOCUMENTOS CANDIDATOS EM: {base.absolute()} ===",
              f"Total: {len(encontrados)} arquivo(s)", ""]

    com_pistas, sem_pistas = [], []
    for rel, path in sorted(encontrados):
        tamanho_kb = max(1, path.stat().st_size // 1024)
        pistas = _detectar_pistas(path)
        registro = (rel.as_posix(), tamanho_kb, pistas)
        (com_pistas if pistas else sem_pistas).append(registro)

    if com_pistas:
        linhas.append("--- PROVAVEIS DOCUMENTOS DE HISTORIAS ---")
        for nome, kb, pistas in sorted(com_pistas, key=lambda r: -len(r[2])):
            linhas.append(f"  {nome}  ({kb} KB)  -> pistas: {', '.join(pistas)}")
        linhas.append("")

    if sem_pistas:
        linhas.append("--- OUTROS DOCUMENTOS (sem pistas de historia) ---")
        for nome, kb, _ in sem_pistas:
            linhas.append(f"  {nome}  ({kb} KB)")
        linhas.append("")

    linhas.append("Use a tool 'ler_documento' com o caminho escolhido para carregar o conteudo.")
    return "\n".join(linhas)


def _detectar_pistas(path: Path) -> list:
    """Le uma amostra do arquivo e retorna as pistas de conteudo de historia."""
    try:
        if path.suffix.lower() in EXT_PDF:
            if pypdf is None:
                return []
            reader = pypdf.PdfReader(str(path))
            amostra = " ".join((p.extract_text() or "") for p in reader.pages[:3])
        elif path.suffix.lower() in EXT_HTML:
            amostra = ler_html(path)[:20_000]
        else:
            amostra = ler_texto_simples(path)[:20_000]
    except Exception:
        return []

    return [descricao for padrao, descricao in PISTAS_HISTORIA if padrao.search(amostra)]


# === Calculo do indice de Definition of Ready ===

PESOS_DOR = {
    "clareza_problema_objetivo": 10,
    "clareza_funcional": 15,
    "criterios_aceite": 15,
    "regras_negocio": 15,
    "dados_validacoes": 10,
    "dependencias_integracoes": 10,
    "seguranca_autorizacao": 5,
    "technical_design": 10,
    "testabilidade": 5,
}

ROTULOS_DOR = {
    "clareza_problema_objetivo": "Clareza do problema e objetivo",
    "clareza_funcional": "Clareza funcional",
    "criterios_aceite": "Criterios de aceite",
    "regras_negocio": "Regras de negocio",
    "dados_validacoes": "Dados e validacoes",
    "dependencias_integracoes": "Dependencias e integracoes",
    "seguranca_autorizacao": "Seguranca e autorizacao",
    "technical_design": "Technical Design",
    "testabilidade": "Testabilidade",
}

# Os pesos listados somam 95%. A pontuacao e normalizada para a escala 0-100
# (indice = pontos x 100 / 95) de modo que notas 100 em todos os criterios
# resultem em indice 100. Se o time recalibrar os pesos para somar 100, a
# normalizacao deixa de ter efeito automaticamente.
SOMA_PESOS_DOR = sum(PESOS_DOR.values())

PESOS_AI = {
    "clareza_determinismo": 25,
    "padroes_conhecidos": 15,
    "complexidade_contida": 15,
    "risco_regressao_baixo": 15,
    "cobertura_testes": 10,
    "exemplos_contratos_estaveis": 10,
    "validacao_automatica": 10,
}

ROTULOS_AI = {
    "clareza_determinismo": "Clareza e determinismo (sem interpretacao)",
    "padroes_conhecidos": "Padroes conhecidos / implementacoes similares",
    "complexidade_contida": "Complexidade contida (poucos arquivos e sistemas)",
    "risco_regressao_baixo": "Risco de regressao baixo",
    "cobertura_testes": "Cobertura de testes existente",
    "exemplos_contratos_estaveis": "Exemplos e estabilidade de contratos",
    "validacao_automatica": "Facilidade de validacao automatica",
}

SOMA_PESOS_AI = sum(PESOS_AI.values())


def _validar_scores(scores: dict, pesos: dict, rotulo: str) -> dict:
    if not isinstance(scores, dict):
        raise ValueError(f"'{rotulo}' deve ser um objeto com os criterios e notas de 0 a 100.")

    desconhecidos = set(scores) - set(pesos)
    if desconhecidos:
        raise ValueError(
            f"Criterio(s) invalido(s) em '{rotulo}': {', '.join(sorted(desconhecidos))}. "
            f"Validos: {', '.join(pesos)}")

    faltantes = set(pesos) - set(scores)
    if faltantes:
        raise ValueError(
            f"Criterio(s) ausente(s) em '{rotulo}': {', '.join(sorted(faltantes))}. "
            "Todos os criterios sao obrigatorios (use 0 quando a informacao nao existe).")

    normalizados = {}
    for chave, valor in scores.items():
        if isinstance(valor, bool) or not isinstance(valor, (int, float)):
            raise ValueError(f"Nota de '{chave}' deve ser numerica entre 0 e 100.")
        if not 0 <= valor <= 100:
            raise ValueError(f"Nota de '{chave}' fora da faixa 0-100: {valor}")
        normalizados[chave] = float(valor)
    return normalizados


def _classificar_dor(indice: float, bloqueadores: int) -> tuple:
    if bloqueadores > 0:
        return ("NOT READY",
                f"Bloqueada: {bloqueadores} bloqueador(es) critico(s). "
                "Um bloqueador critico = NOT READY, independentemente do indice.")
    if indice >= 90:
        return "READY", "Pode construir."
    if indice >= 80:
        return "READY COM RESSALVAS", "Liberada apenas se as ressalvas forem nao-bloqueantes."
    if indice >= 65:
        return "NEEDS REFINEMENT", "Nao iniciar construcao. Refinar antes."
    return "NOT READY", "Bloqueada. Requisitos insuficientes."


def _classificar_ai(indice: float) -> tuple:
    if indice >= 80:
        return "AI-READY", "Execucao por IA recomendada."
    if indice >= 65:
        return "AI-READY-WITH-RESTRICTIONS", "Executar por IA apenas apos os ajustes indicados."
    if indice >= 40:
        return "AI-NEEDS-REFINEMENT", "Execucao por IA nao recomendada no estado atual."
    return "AI-NO", "Seguir fluxo convencional de desenvolvimento."


def _pontuar(scores: dict, pesos: dict, soma_pesos: int) -> tuple:
    """Retorna (pontos brutos, indice normalizado na escala 0-100)."""
    pontos = sum(scores[c] * p / 100 for c, p in pesos.items())
    return pontos, pontos * 100 / soma_pesos


def _tabela_pontuacao(scores: dict, pesos: dict, rotulos: dict,
                      soma_pesos: int, pontos: float) -> list:
    linhas = ["| Criterio | Peso | Nota (0-100) | Pontos |",
              "|---|---|---|---|"]
    for chave, peso in pesos.items():
        nota = scores[chave]
        linhas.append(f"| {rotulos[chave]} | {peso}% | {nota:.0f} | {nota * peso / 100:.2f} |")
    linhas.append(f"| **TOTAL** | **{soma_pesos}%** | - | **{pontos:.2f}** |")
    if soma_pesos != 100:
        linhas.append("")
        linhas.append(f"Os pesos declarados somam {soma_pesos}% (tabela de referencia). "
                      f"A pontuacao e normalizada para a escala 0-100: "
                      f"{pontos:.2f} x 100 / {soma_pesos}.")
    return linhas


def calcular_indice_dor(scores: dict, bloqueadores_criticos: int = 0) -> str:
    normalizados = _validar_scores(scores, PESOS_DOR, "scores")
    if not isinstance(bloqueadores_criticos, int) or bloqueadores_criticos < 0:
        raise ValueError("'bloqueadores_criticos' deve ser um inteiro >= 0.")

    pontos, indice = _pontuar(normalizados, PESOS_DOR, SOMA_PESOS_DOR)
    classificacao, decisao = _classificar_dor(indice, bloqueadores_criticos)

    linhas = ["=== INDICE DE PRONTIDAO - DEFINITION OF READY ===", ""]
    linhas += _tabela_pontuacao(normalizados, PESOS_DOR, ROTULOS_DOR,
                                SOMA_PESOS_DOR, pontos)
    linhas += [
        "",
        f"Indice Definition of Ready: {round(indice)}/100",
        f"Bloqueadores criticos: {bloqueadores_criticos}",
        f"Classificacao: {classificacao}",
        f"Decisao: {decisao}",
    ]

    if bloqueadores_criticos > 0 and indice >= 80:
        linhas.append("")
        linhas.append("ATENCAO: o indice seria suficiente, mas o bloqueador critico "
                      "prevalece e mantem a historia NOT READY.")

    if indice < 80:
        linhas.append("")
        linhas.append("AI-Ready NAO deve ser avaliada: exige Definition of Ready >= 80.")

    return "\n".join(linhas)


def calcular_indice_ai_ready(scores: dict, indice_dor: float) -> str:
    if not isinstance(indice_dor, (int, float)) or isinstance(indice_dor, bool):
        raise ValueError("'indice_dor' deve ser numerico entre 0 e 100.")
    if not 0 <= indice_dor <= 100:
        raise ValueError(f"'indice_dor' fora da faixa 0-100: {indice_dor}")

    if indice_dor < 80:
        return ("=== INDICE DE PRONTIDAO - AI-READY ===\n\n"
                f"Indice Definition of Ready informado: {indice_dor:.0f}/100\n"
                "Resultado: NAO AVALIADA\n"
                "Motivo: somente historias com Definition of Ready >= 80 podem ser "
                "avaliadas para execucao por IA.\n"
                "Proxima acao: refinar a historia e recalcular o Definition of Ready.")

    normalizados = _validar_scores(scores, PESOS_AI, "scores")
    pontos, indice = _pontuar(normalizados, PESOS_AI, SOMA_PESOS_AI)
    classificacao, execucao = _classificar_ai(indice)

    linhas = ["=== INDICE DE PRONTIDAO - AI-READY ===", "",
              "Nota: 'complexidade_contida' e 'risco_regressao_baixo' sao favorabilidade - "
              "quanto MENOR a complexidade/risco, MAIOR a nota.", ""]
    linhas += _tabela_pontuacao(normalizados, PESOS_AI, ROTULOS_AI,
                                SOMA_PESOS_AI, pontos)
    linhas += [
        "",
        f"Indice Definition of Ready: {indice_dor:.0f}/100",
        f"Indice AI-Ready: {round(indice)}/100",
        f"Classificacao: {classificacao}",
        f"Execucao: {execucao}",
    ]

    if classificacao in ("AI-NEEDS-REFINEMENT", "AI-NO") and indice_dor >= 90:
        linhas.append("")
        linhas.append("Observacao: caso valido e esperado - a historia pode estar READY "
                      "para o time e ainda assim nao ser apta para execucao por IA.")

    return "\n".join(linhas)


# === Exportacao ===

def exportar_relatorio_markdown(conteudo: str, caminho: str) -> str:
    path = Path(caminho).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(conteudo, encoding="utf-8")
    return f"Relatorio salvo em: {path.absolute()}"


COLUNAS_VEREDITO = [
    ("id", "ID", 12),
    ("titulo", "Titulo", 45),
    ("definition_of_ready", "Definition of Ready", 24),
    ("indice_dor", "Indice DoR", 12),
    ("ai_ready", "AI-Ready", 28),
    ("indice_ai_ready", "Indice AI", 12),
    ("risco", "Risco", 12),
    ("bloqueadores", "Bloqueadores", 14),
    ("pendencias", "Pendencias", 12),
    ("contradicoes", "Contradicoes", 14),
    ("labels", "Labels / Status", 40),
    ("proxima_acao", "Proxima acao", 40),
]

CORES_DOR = {
    "READY": "C6EFCE",
    "READY COM RESSALVAS": "FFEB9C",
    "NEEDS REFINEMENT": "FFD9A0",
    "NOT READY": "FFC7CE",
}


def exportar_veredito_excel(vereditos: list, caminho: str) -> str:
    if openpyxl is None:
        return ("ERRO: dependencia 'openpyxl' nao instalada. "
                "Execute: pip install -r mcp-server/requirements.txt")
    if not isinstance(vereditos, list) or not vereditos:
        return "ERRO: 'vereditos' deve ser uma lista nao vazia de objetos."

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Vereditos DoR"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="7A1F2B", end_color="7A1F2B", fill_type="solid")

    for col, (_, rotulo, largura) in enumerate(COLUNAS_VEREDITO, 1):
        cell = ws.cell(row=1, column=col, value=rotulo)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = largura

    for row_idx, item in enumerate(vereditos, 2):
        if not isinstance(item, dict):
            continue
        for col, (chave, _, _) in enumerate(COLUNAS_VEREDITO, 1):
            valor = item.get(chave, "")
            if isinstance(valor, (list, tuple)):
                valor = ", ".join(str(v) for v in valor)
            cell = ws.cell(row=row_idx, column=col, value=valor)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

        cor = CORES_DOR.get(str(item.get("definition_of_ready", "")).upper().strip())
        if cor:
            ws.cell(row=row_idx, column=3).fill = PatternFill(
                start_color=cor, end_color=cor, fill_type="solid")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUNAS_VEREDITO)).column_letter}1"

    path = Path(caminho).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return f"Planilha de vereditos salva em: {path.absolute()} ({len(vereditos)} historia(s))"


# === MCP Protocol Handler ===

TOOLS = [
    {
        "name": "ler_documento",
        "description": (
            "Le um documento com historias de usuario e retorna o texto extraido. "
            "Suporta Markdown (.md/.markdown/.txt), HTML (.html/.htm) e PDF (.pdf). "
            "Use esta tool sempre que o usuario anexar ou indicar o caminho de um arquivo "
            "com as historias a serem validadas."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "caminho": {
                    "type": "string",
                    "description": "Caminho (absoluto de preferencia) do arquivo .md, .html ou .pdf"
                }
            },
            "required": ["caminho"]
        }
    },
    {
        "name": "listar_documentos",
        "description": (
            "Varre um diretorio do projeto e lista os arquivos .md, .html, .txt e .pdf, "
            "destacando os que contem pistas de historias de usuario. Use quando o usuario "
            "pedir para validar historias 'do projeto' sem indicar o arquivo exato."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "diretorio": {
                    "type": "string",
                    "description": "Caminho do diretorio a inspecionar (normalmente a raiz do workspace)"
                },
                "profundidade": {
                    "type": "integer",
                    "description": "Profundidade maxima de subpastas a percorrer (padrao 4)"
                }
            },
            "required": ["diretorio"]
        }
    },
    {
        "name": "calcular_indice_dor",
        "description": (
            "Calcula de forma deterministica o Indice de Prontidao (Definition of Ready) "
            "de 0 a 100 aplicando os pesos oficiais e retorna a classificacao "
            "READY / READY COM RESSALVAS / NEEDS REFINEMENT / NOT READY. "
            "Um bloqueador critico forca NOT READY. Informe TODOS os nove criterios com "
            "nota de 0 a 100. Use esta tool em vez de calcular o indice mentalmente."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "scores": {
                    "type": "object",
                    "description": "Notas de 0 a 100 para cada criterio ponderado",
                    "properties": {
                        "clareza_problema_objetivo": {"type": "number", "description": "Peso 10%"},
                        "clareza_funcional": {"type": "number", "description": "Peso 15%"},
                        "criterios_aceite": {"type": "number", "description": "Peso 15%"},
                        "regras_negocio": {"type": "number", "description": "Peso 15%"},
                        "dados_validacoes": {"type": "number", "description": "Peso 10%"},
                        "dependencias_integracoes": {"type": "number", "description": "Peso 10%"},
                        "seguranca_autorizacao": {"type": "number", "description": "Peso 5%"},
                        "technical_design": {"type": "number", "description": "Peso 10%"},
                        "testabilidade": {"type": "number", "description": "Peso 5%"}
                    },
                    "required": [
                        "clareza_problema_objetivo", "clareza_funcional", "criterios_aceite",
                        "regras_negocio", "dados_validacoes", "dependencias_integracoes",
                        "seguranca_autorizacao", "technical_design", "testabilidade"
                    ]
                },
                "bloqueadores_criticos": {
                    "type": "integer",
                    "description": "Quantidade de bloqueadores criticos identificados (padrao 0)"
                }
            },
            "required": ["scores"]
        }
    },
    {
        "name": "calcular_indice_ai_ready",
        "description": (
            "Calcula o Indice AI-Ready de 0 a 100 e retorna a classificacao "
            "AI-READY / AI-READY-WITH-RESTRICTIONS / AI-NEEDS-REFINEMENT / AI-NO. "
            "Exige o indice de Definition of Ready: se for menor que 80, a avaliacao e "
            "recusada por regra. As notas de complexidade e risco sao de favorabilidade "
            "(menor complexidade/risco = nota maior)."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "scores": {
                    "type": "object",
                    "description": "Notas de 0 a 100 para cada criterio de AI-Readiness",
                    "properties": {
                        "clareza_determinismo": {"type": "number", "description": "Peso 25%"},
                        "padroes_conhecidos": {"type": "number", "description": "Peso 15%"},
                        "complexidade_contida": {"type": "number", "description": "Peso 15% (favorabilidade)"},
                        "risco_regressao_baixo": {"type": "number", "description": "Peso 15% (favorabilidade)"},
                        "cobertura_testes": {"type": "number", "description": "Peso 10%"},
                        "exemplos_contratos_estaveis": {"type": "number", "description": "Peso 10%"},
                        "validacao_automatica": {"type": "number", "description": "Peso 10%"}
                    },
                    "required": [
                        "clareza_determinismo", "padroes_conhecidos", "complexidade_contida",
                        "risco_regressao_baixo", "cobertura_testes",
                        "exemplos_contratos_estaveis", "validacao_automatica"
                    ]
                },
                "indice_dor": {
                    "type": "number",
                    "description": "Indice de Definition of Ready ja calculado (0 a 100)"
                }
            },
            "required": ["scores", "indice_dor"]
        }
    },
    {
        "name": "exportar_relatorio_markdown",
        "description": "Salva o relatorio de analise de Definition of Ready em um arquivo Markdown.",
        "inputSchema": {
            "type": "object",
            "properties": {
                "conteudo": {
                    "type": "string",
                    "description": "Conteudo markdown completo do relatorio de analise"
                },
                "caminho": {
                    "type": "string",
                    "description": "Caminho onde salvar o arquivo .md"
                }
            },
            "required": ["conteudo", "caminho"]
        }
    },
    {
        "name": "exportar_veredito_excel",
        "description": (
            "Gera uma planilha Excel consolidada com o veredito de cada historia validada. "
            "Util em validacao em lote de um backlog inteiro."
        ),
        "inputSchema": {
            "type": "object",
            "properties": {
                "vereditos": {
                    "type": "array",
                    "description": "Lista de vereditos, um objeto por historia",
                    "items": {
                        "type": "object",
                        "properties": {
                            "id": {"type": "string"},
                            "titulo": {"type": "string"},
                            "definition_of_ready": {
                                "type": "string",
                                "description": "READY | READY COM RESSALVAS | NEEDS REFINEMENT | NOT READY"
                            },
                            "indice_dor": {"type": "number"},
                            "ai_ready": {
                                "type": "string",
                                "description": "AI-READY | AI-READY-WITH-RESTRICTIONS | AI-NEEDS-REFINEMENT | AI-NO | NAO AVALIADA"
                            },
                            "indice_ai_ready": {"type": "number"},
                            "risco": {"type": "string", "description": "CRITICO | ALTO | MEDIO | BAIXO"},
                            "bloqueadores": {"type": "integer"},
                            "pendencias": {"type": "integer"},
                            "contradicoes": {"type": "integer"},
                            "labels": {"type": "string", "description": "Labels ou status recomendados"},
                            "proxima_acao": {"type": "string"}
                        },
                        "required": ["id", "definition_of_ready"]
                    }
                },
                "caminho": {
                    "type": "string",
                    "description": "Caminho onde salvar o arquivo .xlsx"
                }
            },
            "required": ["vereditos", "caminho"]
        }
    }
]


def _executar_tool(tool_name: str, arguments: dict) -> str:
    if tool_name == "ler_documento":
        return ler_documento(arguments["caminho"])
    if tool_name == "listar_documentos":
        return listar_documentos(arguments["diretorio"],
                                 int(arguments.get("profundidade", 4)))
    if tool_name == "calcular_indice_dor":
        return calcular_indice_dor(arguments.get("scores"),
                                   int(arguments.get("bloqueadores_criticos", 0)))
    if tool_name == "calcular_indice_ai_ready":
        return calcular_indice_ai_ready(arguments.get("scores"),
                                        arguments.get("indice_dor"))
    if tool_name == "exportar_relatorio_markdown":
        return exportar_relatorio_markdown(arguments["conteudo"], arguments["caminho"])
    if tool_name == "exportar_veredito_excel":
        return exportar_veredito_excel(arguments.get("vereditos"), arguments["caminho"])
    return f"Tool desconhecida: {tool_name}"


def handle_request(request: dict):
    """Processa uma requisicao JSON-RPC do MCP."""
    method = request.get("method", "")
    req_id = request.get("id")
    params = request.get("params", {}) or {}

    if method == "initialize":
        return {
            "jsonrpc": "2.0",
            "id": req_id,
            "result": {
                "protocolVersion": "2024-11-05",
                "capabilities": {"tools": {}},
                "serverInfo": {
                    "name": "readiness-mcp",
                    "version": "1.0.0"
                }
            }
        }

    if method == "notifications/initialized":
        return None

    if method == "tools/list":
        return {"jsonrpc": "2.0", "id": req_id, "result": {"tools": TOOLS}}

    if method == "tools/call":
        tool_name = params.get("name", "")
        arguments = params.get("arguments", {}) or {}
        try:
            resultado = _executar_tool(tool_name, arguments)
            return {
                "jsonrpc": "2.0",
                "id": req_id,
                "result": {"content": [{"type": "text", "text": resultado}]}
            }
        except KeyError as exc:
            return {
                "jsonrpc": "2.0",
                "id": req_id,
                "result": {
                    "content": [{"type": "text",
                                 "text": f"ERRO: argumento obrigatorio ausente: {exc}"}],
                    "isError": True
                }
            }
        except Exception as exc:
            return {
                "jsonrpc": "2.0",
                "id": req_id,
                "result": {
                    "content": [{"type": "text", "text": f"ERRO: {exc}"}],
                    "isError": True
                }
            }

    if method == "ping":
        return {"jsonrpc": "2.0", "id": req_id, "result": {}}

    if req_id is not None:
        return {
            "jsonrpc": "2.0",
            "id": req_id,
            "error": {"code": -32601, "message": f"Method not found: {method}"}
        }
    return None


def main():
    """Loop principal - le JSON-RPC via stdin, responde via stdout."""
    sys.stdout.reconfigure(line_buffering=False)

    buffer = ""
    while True:
        try:
            line = sys.stdin.readline()
            if not line:
                break

            buffer += line

            try:
                request = json.loads(buffer.strip())
                buffer = ""
            except json.JSONDecodeError:
                continue

            response = handle_request(request)
            if response is not None:
                sys.stdout.write(json.dumps(response) + "\n")
                sys.stdout.flush()

        except KeyboardInterrupt:
            break
        except Exception as exc:
            sys.stderr.write(f"Erro no servidor: {exc}\n")
            sys.stderr.flush()


if __name__ == "__main__":
    main()
